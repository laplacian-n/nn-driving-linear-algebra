"""Draft 10-minute presentation (16:9) from docs/PRESENTATION.md content.

    python presentation/build_slides.py            -> presentation/slides.pptx  (TH Sarabun New, for PowerPoint)
    PDF_MODE=1 python presentation/build_slides.py -> presentation/slides_pdf.pptx (Sarabun, for the PDF export)
"""
import json
import os

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu, Pt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
PDF = bool(os.environ.get("PDF_MODE"))
FONT = "Sarabun" if PDF else "TH Sarabun New"
SCALE = 0.78 if PDF else 1.0
INK, MUTED, ACCENT, RULE = RGBColor(0x14, 0x14, 0x14), RGBColor(0x55, 0x55, 0x55), RGBColor(0x2a, 0x78, 0xd6), RGBColor(0xdd, 0xdd, 0xdd)
W, H = Emu(12192000), Emu(6858000)          # 13.333 x 7.5 in
CM = 360000


def fig(name):
    return os.path.join(ROOT, "figures", name)


def text_box(slide, x, y, w, h, paras, size=28, color=INK, bold=False):
    tb = slide.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        bullet = p.startswith("- ")
        t = p[2:] if bullet else p
        para.space_after = Pt(6 * SCALE)
        run = para.add_run()
        run.text = ("•  " if bullet else "") + t
        run.font.name, run.font.size, run.font.color.rgb, run.font.bold = FONT, Pt(size * SCALE), color, bold
    return tb


def picture(slide, path, x, y, max_w, max_h):
    w, h = Image.open(path).size
    s = min(max_w / w, max_h / h)
    pw, ph = int(w * s), int(h * s)
    slide.shapes.add_picture(path, Emu(x + (max_w - pw) // 2), Emu(y + (max_h - ph) // 2), Emu(pw), Emu(ph))


def slide(prs, title, bullets=(), image=None, n=None, split=0.42, note=None, stack=False):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    text_box(s, int(0.9 * CM), int(0.6 * CM), int(W - 1.8 * CM), int(1.6 * CM), [title], size=40, bold=True)
    line = s.shapes.add_shape(1, Emu(int(0.9 * CM)), Emu(int(2.25 * CM)), Emu(int(2.2 * CM)), Emu(int(0.12 * CM)))
    line.fill.solid()
    line.fill.fore_color.rgb = ACCENT
    line.line.fill.background()
    top, bottom = int(2.9 * CM), int(H - 1.1 * CM)
    if image and bullets and stack:   # wide figures: short text on top, figure across the slide
        th = int(0.95 * CM) * len(bullets) + int(0.3 * CM)
        text_box(s, int(0.9 * CM), top, int(W - 1.8 * CM), th, list(bullets), size=24)
        picture(s, fig(image), int(0.9 * CM), top + th, int(W - 1.8 * CM), bottom - top - th)
    elif image and bullets:
        tw = int((W - 2.6 * CM) * split)
        text_box(s, int(0.9 * CM), top, tw, bottom - top, list(bullets), size=26)
        picture(s, fig(image), int(0.9 * CM) + tw + int(0.8 * CM), top, int(W - 1.8 * CM) - tw - int(0.8 * CM), bottom - top)
    elif image:
        picture(s, fig(image), int(0.9 * CM), top, int(W - 1.8 * CM), bottom - top)
    else:
        text_box(s, int(0.9 * CM), top, int(W - 1.8 * CM), bottom - top, list(bullets), size=28)
    if n:
        text_box(s, int(W - 2.2 * CM), int(H - 1.0 * CM), int(1.6 * CM), int(0.8 * CM), [str(n)], size=16, color=MUTED)
    if note:
        s.notes_slide.notes_text_frame.text = note
    return s


def showcase_numbers():
    d = os.path.join(ROOT, "experiments", "results", "showcase")
    best = None
    for seed in sorted(os.listdir(d)) if os.path.exists(d) else []:
        p = os.path.join(d, seed, "final.json")
        if os.path.exists(p):
            r = json.load(open(p))
            if best is None or r["bc_ga"]["show"]["score"] > best["bc_ga"]["show"]["score"]:
                best = r
    if not best:
        return ["- ผลกำลังรอการฝึกเสร็จ"]
    b, g = best["bc_base"], best["bc_ga"]
    ls = lambda x: f"{x:.1f}" if x else "–"
    return [f"- สนามโหดระดับ 3.0 ที่ไม่เคยเห็น: จบรอบ {100 * b['show']['finish']:.0f}% → {100 * g['show']['finish']:.0f}%",
            f"- ความเร็วรอบที่จบ {ls(b['show']['lap_speed'])} → {ls(g['show']['lap_speed'])} ม./วิ. "
            f"(ครู {sum(best['teacher_show_v_ref']) / len(best['teacher_show_v_ref']):.1f}) เร็วกว่าครูแล้ว",
            f"- แต่ลืมสนามปกติบางส่วน: ชุด T {100 * b['T']['finish']:.0f}% → {100 * g['T']['finish']:.0f}%"]


def main():
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    n = 0

    def nxt():
        nonlocal n
        n += 1
        return n

    s = prs.slides.add_slide(prs.slide_layouts[6])
    picture(s, fig("screens/app_bc_xl_d20.png"), int(W * 0.45), 0, int(W * 0.55), int(H))
    text_box(s, int(0.9 * CM), int(2.2 * CM), int(W * 0.42), int(5 * CM),
             ["2D Map", "Self-Driving Car", "Neural Network"], size=40, bold=True)
    text_box(s, int(0.9 * CM), int(8.4 * CM), int(W * 0.42), int(6 * CM),
             ["รถแข่งขับเคลื่อนอัตโนมัติในแผนที่สองมิติด้วยโครงข่ายประสาทเทียม", "", "รายวิชา Linear Algebra",
              "ผู้จัดทำ ____________________"], size=26, color=MUTED)
    nxt()

    slide(prs, "ที่มาของปัญหา", [
        "- รถขับเองต้องตัดสินใจจากเซนเซอร์หลายมิติทุกเสี้ยววินาที",
        "- ต้องคุมพวงมาลัย คันเร่ง เบรก พร้อมกัน โดยเฉพาะตอนยางใกล้หลุด",
        "- เขียนกฎให้ครบทุกสถานการณ์ทำได้ยาก",
        "- คำถาม: สมองต้องใหญ่แค่ไหน? เซนเซอร์กี่ตัว? ฝึกแบบไหน? บีบอัดได้แค่ไหน?",
    ], "tracks_d10.png", n=nxt())

    slide(prs, "ไอเดียและแนวทางที่เลือก", [
        "- จำลองรถแข่ง 2D ฟิสิกส์สมจริง + สนาม F1 สุ่มใหม่ทุกครั้ง (ความยาก 0–3)",
        "- รถเห็นสนามผ่านรังสีวัดระยะ 9 เส้นเท่านั้น",
        "- neural network: 13 อินพุต → 4 คำสั่ง",
        "- ฝึก 2 แบบ: Genetic Algorithm (เรียนเอง) และเรียนจากครูด้วย gradient",
        "- เขียนเองทั้งหมดด้วย NumPy",
    ], "screens/app_ga_s_d10.png", n=nxt())

    slide(prs, "ภาพรวมขั้นตอนและการเตรียมข้อมูล", [
        "- สนาม → ฟิสิกส์ → เซนเซอร์ → เวกเตอร์ 13 มิติ → neural network → คำสั่ง",
        "- ข้อมูลการขับ 22,819 แถว: สมบูรณ์ 77.8% / แก้ไข 16.8% / ตัดทิ้ง 5.4%",
        "- หลังทำความสะอาด PCA 8 แกนอธิบายได้ 90.8% (ก่อนหน้า 78.8%)",
        "- ฝึกตามหลักสูตร 7 สถานี แล้วสอบ 210 สนามที่ไม่เคยเห็น",
    ], "data_pca_scree.png", n=nxt())

    slide(prs, "พีชคณิตเชิงเส้นที่ใช้", [
        "- เวกเตอร์: การเคลื่อนที่ (บวก + หมุน), ผลคูณจุดหาระยะบนสนามและการชน, ผลคูณไขว้หาความโค้ง, การกลายพันธุ์ = บวกเวกเตอร์",
        "- เมทริกซ์: กฎของคราเมอร์หาจุดตัดรังสี, คูณเมทริกซ์ใน neural network, ค่าไอเกน (PCA), ทรานสโพส/ผลคูณภายนอก (backprop), สมการปกติ (ตัดโหนด), SVD",
    ], "data_eigen_sensors.png", n=nxt(), split=0.45)

    slide(prs, "ผล 1: Genetic Algorithm กับขนาดโมเดล", [
        "- ระหว่างฝึกทุกขนาดได้ 74–78% เท่ากัน",
        "- หลังฝึก S 2×8 ดีสุด 81%, XL 60%",
        "- โมเดลใหญ่ลืมสนามง่าย (XL ระดับ 0.5: 90 → 58%)",
        "- ไม่ใช่เพราะเกณฑ์อิ่มตัว → GA ค้นหาในมิติสูงได้ยาก",
    ], "s3_overall.png", n=nxt())

    slide(prs, "ผล 2: เรียนจากครูด้วย gradient", [
        "- โมเดลใหญ่ชนะจริง (ครู 75%): XS 16 → S 82 → L 88 → XL 94%   •   ฝึกเร็วกว่า GA ราว 10 เท่า",
        "- งานยิ่งต้องแม่น (ครูขับเร็ว) โมเดลใหญ่ยิ่งได้เปรียบ   •   ครูขับเผื่อระยะสอนได้ดีกว่าครูขับสุดขีด",
    ], "cmp_variants_by_level.png", n=nxt(), stack=True)

    slide(prs, "ผล 3: เซนเซอร์และ PCA", [
        "- รังสี 3 เส้น 0%, 9 เส้น 88%, 17 เส้น 96%   •   PCA 4 มิติจาก 17 รังสี = 93% ดีกว่ารังสีจริง 9 เส้น",
        "- ทิศสุ่ม 4 มิติได้แค่ 50% → เวกเตอร์ไอเกนมีความหมาย   •   PCA 17 มิติ = หมุนแกน = ผลเท่าเดิม",
    ], "sens_dimensions.png", n=nxt(), stack=True)

    slide(prs, "ผล 4: การบีบอัดโมเดล", [
        "- 8 bit: เล็กลง 3.6 เท่า แทบไม่เสีย (94 → 92%)",
        "- XL 8 bit ดีกว่า L ที่ฝึกเอง ทั้งที่เล็กกว่า",
        "- ตัดโหนด / low-rank SVD พัง: ความคลาดเคลื่อนสะสมในการขับแบบวงปิด",
    ], "s4bc_pareto_XL.png", n=nxt())

    slide(prs, "โมเดลตัวโชว์: สนามโหดระดับ 3.0", [
        "- XL + เรียนจากครู → เรียนต่อบนสนามถนนกว้าง 7 เมตร → GA ฝึกต่อเอง",
        *showcase_numbers(),
    ], "tracks_d30.png", n=nxt())

    demo_fig = "demo_results.png" if os.path.exists(fig("demo_results.png")) else "screens/demo_menu.png"
    slide(prs, "Demo: ให้เพื่อนทดสอบ", [
        "- เพื่อน 2 คน เลือกสนาม (สุ่ม / seed / วาดเอง) ความยาก และโมเดลเอง",
        "- รวม 10+ ครั้ง ผลบันทึกอัตโนมัติ",
        "- สรุปผลรายคนและรายโมเดล",
    ], demo_fig, n=nxt())

    slide(prs, "สรุปและปัญหาที่พบ", [
        "- ขนาดโมเดลที่เหมาะขึ้นกับวิธีฝึก • ทบทวนข้อมูลเก่ากันการลืม • ครูที่ดีสำคัญกว่าขนาด • PCA ลดมิติได้มาก • 8 bit บีบอัดได้ปลอดภัย",
        "- ปัญหา: รถหมุนง่ายเกินจริง → ปรับยาง, เซนเซอร์ช้า → Numba 8×, ลืมถนนกว้าง → rehearsal, เวลาไม่พอ → ลดขนาดการทดลอง + BC, ตัดโหนดพัง → 8 bit",
        "- ต่อยอด: reinforcement learning ด้วย gradient, จำลองบน GPU, เพิ่ม seed",
        "- โค้ด: github.com/laplacian-n/nn-driving-linear-algebra",
    ], n=nxt())

    out = os.path.join(ROOT, "presentation", "slides_pdf.pptx" if PDF else "slides.pptx")
    prs.save(out)
    print("wrote", out, n, "slides")


if __name__ == "__main__":
    main()
